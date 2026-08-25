#!/usr/bin/env python3
"""
colt_importar_csv.py
=====================
Vuelca 04_BASE_FINAL/base_colt_sudamerica_trabajo.csv (el CSV de trabajo, ya con las
correcciones/filas nuevas de la campaña aplicadas) de vuelta a la hoja
"Datos_COLT_Sudamerica" de Base_COLT_Sudamerica.xlsx. Es el ultimo paso de una
campaña (ver colt_export_csv.py para el flujo completo).

Como evita el bug de performance del 2026-08-24 (Uruguay, esquema anterior): en vez de
editar celda por celda con ws.cell(row=, column=).value = ... sobre un workbook
read_only (carisimo, O(n) por celda), este script:
  1. Abre el xlsx en modo NORMAL (no read_only) -necesario para poder borrar/escribir filas-.
  2. Borra TODAS las filas de datos existentes de un saque (ws.delete_rows), conserva
     encabezado y formato de la hoja.
  3. Reescribe las filas con ws.append(), que es secuencial y rapido.
  4. Guarda. Las otras hojas (Leame, Viajes_Cancelados) NO se tocan.

Antes de correr esto, correr colt_validar_madre.py y revisar que de 0 errores -este
script no valida nada, solo vuelca lo que encuentra en el csv-.

USO
---
    python colt_importar_csv.py
"""
import csv
from datetime import date
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
XLSX_PATH = PROJECT_ROOT / "04_BASE_FINAL" / "Base_COLT_Sudamerica.xlsx"
SHEET_NAME = "Datos_COLT_Sudamerica"
CSV_TRABAJO = PROJECT_ROOT / "04_BASE_FINAL" / "base_colt_sudamerica_trabajo.csv"

CAMPOS_FECHA = {"TripStartDate", "TripEndDate"}
CAMPOS_ENTEROS = {"TripYear", "TripDuration"}


def parsear_valor(campo, valor):
    if valor is None or valor == "":
        return None
    if campo in CAMPOS_FECHA:
        try:
            y, m, d = (int(x) for x in str(valor)[:10].split("-"))
            return date(y, m, d)
        except Exception:
            return valor  # texto raro (fecha estimada, etc.) se guarda tal cual
    if campo in CAMPOS_ENTEROS:
        try:
            return int(float(valor))
        except (ValueError, TypeError):
            return valor
    return valor


def main():
    if not CSV_TRABAJO.exists():
        raise SystemExit(f"No existe {CSV_TRABAJO}. Correr colt_export_csv.py primero.")

    with open(CSV_TRABAJO, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        headers = next(reader)
        filas = list(reader)

    # Chequeo de seguridad: nunca volcar un csv mas chico que un umbral razonable
    # (evita pisar la hoja madre con un csv vacio o cortado por error).
    if len(filas) < 1000:
        raise SystemExit(
            f"ABORTADO: el csv de trabajo tiene solo {len(filas)} filas -sospechosamente "
            f"pocas para la base completa de Sudamerica (deberian ser miles)-. Revisar "
            f"antes de forzar el volcado."
        )

    wb = openpyxl.load_workbook(XLSX_PATH)  # NORMAL, no read_only -hace falta para escribir
    ws = wb[SHEET_NAME]

    encabezado_actual = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if encabezado_actual != headers:
        raise SystemExit(
            "ABORTADO: el encabezado del csv de trabajo no coincide exactamente con el de "
            "la hoja Datos_COLT_Sudamerica. No se toca el xlsx. Revisar manualmente."
        )

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    idx_por_campo = {h: i for i, h in enumerate(headers)}
    for fila in filas:
        fila_tipada = [
            parsear_valor(h, fila[idx_por_campo[h]]) for h in headers
        ]
        ws.append(fila_tipada)

    wb.save(XLSX_PATH)
    print(f"Volcado: {len(filas)} filas -> {XLSX_PATH} (hoja {SHEET_NAME})")


if __name__ == "__main__":
    main()
